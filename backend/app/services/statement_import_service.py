"""Parsers for provider files. Files are processed in memory and never persisted."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import re

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.account import Account
from app.models.enums import ExchangeRateSource, TransactionPurpose, TransactionType
from app.models.transaction import Transaction
from app.schemas.statement_import import StatementCommit, StatementCommitResult, StatementPreview, StatementRow
from app.services.exchange_rate_service import get_exchange_rate

TRADERNET_HEADERS = ("Операция №", "Дата", "Операция", "Комментарий", "Сумма", "Валюта")


def _symbol(comment: str) -> str | None:
    match = re.search(r"\(([A-Z0-9][A-Z0-9.]{1,40})\)\)", comment)
    return match.group(1) if match else None


def _event_kind(operation: str, amount: Decimal) -> tuple[TransactionType, TransactionPurpose, bool, str | None]:
    normalized = operation.strip().casefold()
    if normalized == "дивиденды":
        return TransactionType.INCOME, TransactionPurpose.DIVIDEND, True, None
    if normalized == "купон":
        return TransactionType.INCOME, TransactionPurpose.COUPON, True, None
    if normalized == "комиссия за сделки":
        return TransactionType.EXPENSE, TransactionPurpose.FEE, True, None
    if normalized == "налоги":
        return TransactionType.EXPENSE, TransactionPurpose.TAX, True, None
    if normalized == "карточный платеж":
        return TransactionType.EXPENSE, TransactionPurpose.ORDINARY, True, None
    if normalized in {"блокировка", "разблокировка"}:
        return TransactionType.EXPENSE, TransactionPurpose.ORDINARY, False, "Temporary reservation is not a posted cash movement"
    if normalized == "оплата по сделке":
        return TransactionType.EXPENSE, TransactionPurpose.INVESTMENT_TRADE, False, "Trade settlement needs the broker trades report to avoid duplicate principal"
    if normalized == "перевод внутри компании":
        return TransactionType.TRANSFER, TransactionPurpose.ORDINARY, False, "Internal transfer needs both source and destination brokerage accounts"
    return (TransactionType.INCOME if amount >= 0 else TransactionType.EXPENSE), TransactionPurpose.ORDINARY, True, None


def parse_tradernet_xlsx(content: bytes, file_name: str) -> StatementPreview:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(422, "The XLSX file could not be opened") from exc
    rows: list[StatementRow] = []
    warnings: list[str] = []
    matched_sheet = False
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header or tuple(str(value).strip() if value is not None else "" for value in header[:6]) != TRADERNET_HEADERS:
            warnings.append(f"{sheet.title}: unsupported header row")
            continue
        matched_sheet = True
        for row_number, values in enumerate(iterator, start=2):
            if not any(value is not None for value in values):
                continue
            try:
                operation_id, raw_date, operation, comment, raw_amount, currency = values[:6]
                event_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
                if not isinstance(event_date, date):
                    event_date = datetime.strptime(str(raw_date), "%d.%m.%Y").date()
                amount = Decimal(str(raw_amount))
                if amount == 0:
                    raise ValueError("zero amount")
                operation_text = str(operation).strip()
                tx_type, purpose, importable, warning = _event_kind(operation_text, amount)
                external_id = f"tradernet:{str(operation_id).strip()}"
                rows.append(StatementRow(
                    source_row=f"{sheet.title}!{row_number}", external_id=external_id, date=event_date,
                    type=tx_type, amount=abs(amount), currency=str(currency).strip().upper(),
                    description=operation_text, details=str(comment).strip() if comment else None,
                    purpose=purpose, security_symbol=_symbol(str(comment or "")), importable=importable, warning=warning,
                ))
            except (ValueError, InvalidOperation, TypeError) as exc:
                warnings.append(f"{sheet.title}!{row_number}: {exc}")
    if not matched_sheet:
        raise HTTPException(422, "No Tradernet cash-movement sheet was recognized")
    return StatementPreview(provider="tradernet", file_name=file_name, rows=rows, warnings=warnings)


async def commit_statement(session: AsyncSession, payload: StatementCommit) -> StatementCommitResult:
    importable = [row for row in payload.rows if row.importable]
    ids = [row.external_id for row in importable]
    existing = set((await session.execute(select(Transaction.external_id).where(Transaction.external_id.in_(ids)))).scalars().all())
    accounts: dict[str, Account] = {}
    for currency, account_id in payload.accounts_by_currency.items():
        account = await session.get(Account, account_id)
        if account is None:
            raise HTTPException(400, f"Unknown account for {currency}")
        if account.currency.upper() != currency.upper():
            raise HTTPException(422, f"Account {account.name} is not denominated in {currency.upper()}")
        accounts[currency.upper()] = account
    created = 0
    for row in importable:
        if row.external_id in existing:
            continue
        account = accounts.get(row.currency)
        if account is None:
            raise HTTPException(422, f"No destination account selected for {row.currency}")
        if row.currency == "KZT":
            rate, source = Decimal("1"), ExchangeRateSource.NBK
        else:
            try:
                official = await get_exchange_rate(session, row.date, row.currency)
                rate, source = official.rate_to_kzt, ExchangeRateSource.NBK
            except HTTPException as exc:
                if exc.status_code not in {422, 503}:
                    raise
                rate, source = None, None
        session.add(Transaction(
            account_id=account.id, category_id=None, type=row.type, amount=row.amount,
            exchange_rate_to_kzt=rate,
            base_amount_kzt=(row.amount * rate).quantize(Decimal("0.01")) if rate is not None else None,
            exchange_rate_source=source, external_id=row.external_id, purpose=row.purpose,
            description=row.description, notes=row.details, date=row.date,
        ))
        created += 1
    await session.commit()
    return StatementCommitResult(created=created, duplicates=len(existing.intersection(ids)), ignored=len(payload.rows) - len(importable))
